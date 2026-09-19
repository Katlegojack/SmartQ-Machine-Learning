from docx import Document
from pathlib import Path

PATH = Path("docs/SmartQ_100k_Synthetic_Dataset_Documentation.docx")
doc = Document(PATH)

replacements = {
    6: "I prepared this document as supporting documentation for my SmartQ synthetic dataset.",
    9: "In this document I explain what my 100,000-row SmartQ dataset represents, how I generated the synthetic records, what the fields mean, how I validated the dataset, and how I use it in my statistical and machine-learning work.",
    13: "I generated my SmartQ Synthetic Operational Dataset with 100,000 queue-related records using a reproducible simulation of the SmartQ operating model. I treat each row as one customer visit or scheduled visit. The row records information available before service, the queue conditions around the visit, and - when the visit is completed - the observed waiting-time and service outcomes.",
    18: "I allow daily counter availability to vary slightly so the simulation can represent ordinary staffing or maintenance effects. This changes queue capacity and therefore changes waiting-time outcomes.",
    21: "I deliberately allow synthetic service duration to finish earlier than target, near target, or later than target. I also include a small long-tail component so some services take much longer. I did this because a fixed service time for every customer would create an unrealistically easy prediction problem.",
    23: "I designed SmartQ to collect operational observations during real queue activity, but at this stage I do not yet have enough real labelled production data for a large ML experiment. I therefore generated a synthetic dataset that mirrors the type of operational evidence SmartQ could accumulate over many operating days.",
    24: "I used synthetic generation for three practical reasons:",
    25: "Scale: I can work with 100,000 observations immediately for model development and evaluation.",
    26: "Reproducibility: I used a fixed random seed so I can regenerate the same synthetic design consistently.",
    27: "Privacy: I did not include real names, usernames, email addresses, phone numbers, dates of birth, disability information or pregnancy information.",
    30: "I use this dataset to demonstrate data preparation, EDA, baseline construction, model training, validation and error comparison. I do not use it to claim that my trained model will achieve the same accuracy in a real production queue. Before real deployment, I would still need ethically collected operational data and external validation.",
    32: "I did not generate independent random spreadsheet rows. I simulated operating days so queue conditions and outcomes are connected. When demand is high, counters become busy, customers accumulate, people_ahead rises and waits can grow. When demand drops or more counter capacity is available, queues can clear faster.",
    34: "I generate operating activity on weekdays. I place appointments in 15-minute slots between 08:00 and 17:45. I deliberately make demand non-uniform so morning and afternoon periods are busier while lunch and late-day periods are usually quieter. I also vary demand by branch so branch_code contains operational information.",
    36: "I generate about 78% of visits as appointments and the rest as walk-ins. Appointment customers can arrive early, on time or late. Walk-ins receive an arrival time directly instead of an appointment time.",
    39: "I keep physical arrival/check-in separate from service eligibility. An appointment customer can arrive and check in early, but I do not make that customer call-eligible before the booked appointment time. A walk-in becomes eligible at check-in. I store this distinction explicitly in service_eligible_at.",
    42: "For attended visits, I calculate SmartQ-style queue features around check-in time. These include people ahead in the same lane, General/Priority waiting counts, customers being served, open counter capacity, counter utilisation, queue pressure, estimated workload ahead and recent throughput.",
    43: "I calculate rolling-history features using only visits whose outcomes were already known before the current prediction time. For example, recent_avg_service_minutes_10 uses the ten most recently completed services known at that customer's check-in. I did this to reduce look-ahead bias.",
    45: "In this version of my synthetic generator, I route General customers to General counters and Priority customers to Priority counters. Inside each eligible lane I process customers in first-come-first-served order, subject to counter availability. I also give each counter a small daily efficiency variation and let service duration depend on service type and a late-day fatigue factor.",
    47: "I use baseline_eta_minutes as my deterministic SmartQ-style estimate before the real wait outcome is known. I combine the delay until service eligibility with estimated workload ahead and effective open capacity. I do not treat this as an ML prediction; I keep it as an engineering benchmark for the learned models.",
    49: "I represent each synthetic visit using this lifecycle:",
    50: "I create a visit as an appointment or walk-in.",
    51: "For an appointment, I allow the customer to cancel, no-show, or arrive and check in.",
    52: "When a customer checks in, I calculate a service-eligibility time.",
    53: "At the prediction/check-in point, I capture queue and capacity conditions.",
    54: "When a matching counter becomes available, I call the customer.",
    55: "I start service after a short handoff delay.",
    56: "I end service after the generated actual service duration.",
    57: "After that, I have labelled outcomes such as actual wait, service duration and variance values for analysis.",
    61: "I calculated the following values directly from my generated 100,000-row file. They are measured properties of the synthetic dataset, not hypothetical targets.",
    71: "My dataset contains 45 columns. I use “Pre-outcome” for fields that can be available at or before the prediction point. I use “Outcome” for fields observed only after the queue/service process has progressed. I generally exclude outcome fields when I predict waiting time at check-in.",
    75: "My main supervised-learning task is wait-time regression. For completed visits, I use actual_wait_minutes as the target. I give the model information that would have been available at check-in and ask it to predict how long that customer waits before being called.",
    76: "7.2 Pre-outcome predictors I considered",
    77: "I considered branch_code, service_code, booking_source and queue_type.",
    78: "I considered arrival_offset_minutes and time features derived from service eligibility.",
    79: "I considered people_ahead, general_waiting, priority_waiting and serving_count.",
    80: "I considered open_general_counters, open_priority_counters and effective_open_counters.",
    81: "I considered counter_utilisation, queue_pressure_index and workload_minutes_ahead.",
    82: "I considered recent_avg_service_minutes_10, recent_avg_wait_minutes_10 and recent_throughput_60m.",
    83: "I considered service_target_minutes, hour_of_day, day_of_week and is_peak_period.",
    84: "I kept baseline_eta_minutes as a comparison benchmark rather than an official model input in my final experiment.",
    85: "7.3 Leakage fields I exclude from wait-time features",
    86: "When I predict wait at check-in, I do not allow the model to receive information that only becomes known after the customer is called or served. I exclude at least the following fields:",
    98: "Because my data is time ordered, I prefer a chronological split to randomly shuffling all rows. In my final experiment I split by whole operating dates so earlier dates are used for training, later dates for validation, and the latest dates for testing.",
    100: "I use MAE and RMSE as my main regression metrics, and I also inspect R², residual behaviour and other diagnostics. I compare my learned models against both a simple mean-wait baseline and the existing deterministic baseline_eta_minutes benchmark.",
    105: "My dataset is useful for research and learning, but I keep its boundaries explicit in my report and presentation.",
    111: "10. My Final Summary",
    112: "I built the SmartQ 100,000-row synthetic dataset as a privacy-preserving research dataset that represents queue demand, appointment and walk-in behaviour, service eligibility, queue pressure, counter capacity, baseline ETA, waiting outcomes and service outcomes over many simulated operating days. I linked the features temporally instead of creating independent random rows so I could run more realistic queue-analytics and waiting-time experiments.",
    113: "For my Special Topic project, I use this dataset as a controlled evidence base for comparing waiting-time prediction approaches. I keep the synthetic nature explicit, prevent target leakage, split data chronologically, compare learned models with simple and deterministic baselines, and reserve real-world deployment claims until I have genuine SmartQ operational data.",
}

for index, text in replacements.items():
    doc.paragraphs[index].text = text

callouts = {
    2: "Important academic disclosure\nI generated this dataset synthetically. I designed it to reproduce queue behaviour and ML conditions, but I do not present it as real customer data.",
    6: "My priority-privacy design\nI record only the resulting queue_type (GENERAL or PRIORITY). I do not include sensitive reasons such as disability or pregnancy in the synthetic ML dataset.",
    8: "Why this matters to me\nBy separating check-in from service eligibility, I prevent my simulation from serving appointment customers before their booked time. This fixes one of the important lessons I learned from earlier simulation work.",
    11: "How I interpret the long wait tail\nI intentionally allow the maximum wait to be much larger than the median because severe congestion is part of the problem SmartQ is meant to understand. I do not remove a long wait simply because it makes the model harder to fit.",
    18: "My model-selection principle\nI only adopt a more complex model when it improves validation performance enough to justify the extra complexity. In my final project comparison I select the official model using validation MAE, not final test results.",
    20: "My recommended disclosure\nI generated a 100,000-row synthetic operational dataset to support SmartQ ML development. I use it to demonstrate the modelling process, not to claim guaranteed production accuracy.",
}

for table_index, text in callouts.items():
    doc.tables[table_index].cell(0, 0).text = text

split_table = doc.tables[16]
split_table.cell(0, 1).text = "My actual rule"
split_table.cell(0, 2).text = "Why I use it"
split_rows = [
    ("Training", "2026-01-02 to 2026-06-16", "I fit preprocessing and model parameters."),
    ("Validation", "2026-06-17 to 2026-07-22", "I compare models and limited tuning choices."),
    ("Test", "2026-07-23 to 2026-08-27", "I use this as my final later-date evaluation."),
]
for row_index, row in enumerate(split_rows, start=1):
    for column_index, value in enumerate(row):
        split_table.cell(row_index, column_index).text = value

model_table = doc.tables[17]
for row in range(1, len(model_table.rows)):
    for column in range(len(model_table.columns)):
        model_table.cell(row, column).text = ""

model_rows = [
    ("Linear Regression", "I use it as my simple, interpretable regression reference."),
    ("Random Forest", "I use it to capture nonlinear queue interactions with many averaged trees."),
    ("XGBoost", "I use it as a strong boosted-tree model for structured SmartQ data."),
    ("Future hybrid approach", "I could later test ML that learns residual error on top of my deterministic ETA, but I kept this outside the official comparison."),
]
for row_index, row in enumerate(model_rows, start=1):
    for column_index, value in enumerate(row):
        model_table.cell(row_index, column_index).text = value

files_table = doc.tables[21]
for row in range(1, len(files_table.rows)):
    purpose = files_table.cell(row, 1).text
    if "Full 100,000-row" in purpose:
        files_table.cell(row, 1).text = "I use this as the full 100,000-row dataset for analysis and machine learning."
    elif "Read-me" in purpose:
        files_table.cell(row, 1).text = "I use this as a companion read-me, validation summary, data dictionary and 1,000-row preview."
    elif "narrative documentation" in purpose:
        files_table.cell(row, 1).text = "I use this document to explain how I designed, validated and use the dataset."

doc.save(PATH)
print(f"Rewrote {PATH} in first-person voice.")
